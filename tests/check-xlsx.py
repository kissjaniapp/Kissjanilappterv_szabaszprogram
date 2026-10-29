"""Validate the actual exported .xlsx using the Excel-compatible openpyxl reader."""
import sys
from openpyxl import load_workbook
path = sys.argv[1]
wb = load_workbook(path, read_only=False, data_only=True)
ws = wb.active
assert len(wb.sheetnames) == 1, wb.sheetnames
assert ws["A1"].value == "ANYAG: Feya", ws["A1"].value
assert ws["A2"].value == "HOVÁ KÉSZÜL: Kiskunfélegyháza", ws["A2"].value
assert ws["A4"].value == "Alkatrész megnevezése", ws["A4"].value
assert [ws.cell(4,i).value for i in range(1,6)] == [
    "Alkatrész megnevezése", "Kész X (mm)", "Kész Y (mm)",
    "Darab", "Élzárás"
]
assert [ws.cell(5,i).value for i in range(1,6)] == [
    "Oldallap", 910, 600, 2, "1-0 · 2 mm · 1 mm előmarás"
]
assert ws.max_column == 5, ws.max_column
print("Excel xlsx megnyithato es az 5 oszlop hibamentes:", path)
